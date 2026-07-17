"""Tests for modules/collections_store.py — run with:
.venv/bin/python tests/test_collections.py

Uses a temporary collections folder (with a space in the path on purpose).
Extraction tests are model-free; index/search tests load Qwen3-Embedding 0.6B.
OCR tests use Apple Vision on generated fixture images.
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from modules import storage  # noqa: E402

_TMP = Path(tempfile.mkdtemp(prefix="coll test "))  # space in path on purpose
storage.collections_dir = lambda: _TMP  # keep tests away from real user data

from modules import collections_store as cs  # noqa: E402

FAILURES = []


def check(name, cond, detail=""):
    if cond:
        print(f"  ok  {name}")
    else:
        print(f"FAIL  {name}  {detail}")
        FAILURES.append(name)


def make_fixtures(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    files = {}

    files["txt"] = d / "hsk words.txt"
    files["txt"].write_text(
        "爱 ài — to love; tình yêu. HSK level 1.\n"
        "喜欢 xǐhuan — to like; thích. HSK level 1.\n"
        "学习 xuéxí — to study; học tập. HSK level 1.\n",
        encoding="utf-8",
    )

    files["csv"] = d / "hsk_list.csv"
    files["csv"].write_text(
        "hanzi,pinyin,hsk,meaning_en,meaning_vi\n"
        "爱,ài,1,to love,yêu\n"
        "喜欢,xǐhuan,1,to like,thích\n"
        "学习,xuéxí,1,to study,học\n",
        encoding="utf-8",
    )

    files["json"] = d / "dict.json"
    files["json"].write_text(json.dumps({
        "words": [
            {"hanzi": "爱", "traditional": "愛", "pinyin": "ài", "vi": "yêu"},
            {"hanzi": "喜欢", "traditional": "喜歡", "pinyin": "xǐhuan", "vi": "thích"},
        ]
    }, ensure_ascii=False), encoding="utf-8")

    files["md"] = d / "grammar.md"
    files["md"].write_text(
        "# 爱 as a verb\n爱 + noun expresses love for something. 我爱你。\n\n"
        "# 喜欢 usage\n喜欢 is milder than 爱. 我喜欢学习。\n",
        encoding="utf-8",
    )

    import docx
    doc = docx.Document()
    doc.add_heading("Example sentences 例句", level=1)
    doc.add_paragraph("我爱我的家。 Wǒ ài wǒ de jiā. I love my family.")
    doc.add_paragraph("他喜欢喝茶。 Tā xǐhuan hē chá. He likes drinking tea.")
    files["docx"] = d / "examples.docx"
    doc.save(str(files["docx"]))

    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype("/System/Library/Fonts/STHeiti Light.ttc", 48)
    img = Image.new("RGB", (900, 220), "white")
    draw = ImageDraw.Draw(img)
    draw.text((30, 30), "爱 ai — to love — yêu", font=font, fill="black")
    draw.text((30, 110), "HSK level 1 vocabulary", font=font, fill="black")
    files["png"] = d / "scan note.png"
    img.save(files["png"])

    img2 = Image.new("RGB", (900, 220), "white")
    draw2 = ImageDraw.Draw(img2)
    draw2.text((30, 30), "学习 xuexi — to study — học tập", font=font, fill="black")
    draw2.text((30, 110), "Scanned page for OCR test", font=font, fill="black")
    files["pdf_scanned"] = d / "scanned.pdf"
    img2.convert("RGB").save(files["pdf_scanned"], "PDF")

    return files


def test_extraction(files):
    recs = cs.extract_records(files["txt"])
    check("txt extracted", recs and "爱" in recs[0]["text"])

    recs = cs.extract_records(files["csv"])
    check("csv rows with numbers", recs and recs[0].get("row") and "hanzi: 爱" in recs[0]["text"])

    recs = cs.extract_records(files["json"])
    check("json paths kept", recs and recs[0].get("json_path", "").startswith("$")
          and any("愛" in r["text"] for r in recs))

    recs = cs.extract_records(files["md"])
    check("md sections", any(r.get("section") == "爱 as a verb" for r in recs))

    recs = cs.extract_records(files["docx"])
    check("docx section heading", any("例句" in (r.get("section") or "") for r in recs))

    recs = cs.extract_records(files["png"])
    check("image OCR text", recs and "HSK" in recs[0]["text"], recs)

    recs = cs.extract_records(files["pdf_scanned"])
    check("scanned pdf OCR fallback",
          recs and recs[0].get("section") == "OCR" and recs[0].get("page") == 1
          and ("study" in recs[0]["text"] or "OCR" in recs[0]["text"]), recs)


def test_collections_crud(files):
    meta = cs.create_collection("HSK word lists", "official lists")
    check("create", meta["name"] == "HSK word lists")
    try:
        cs.create_collection("HSK word lists")
        check("duplicate create rejected", False)
    except ValueError:
        check("duplicate create rejected", True)

    msg = cs.add_files("HSK word lists", [str(files["txt"]), str(files["csv"])])
    check("add files", "2 file(s)" in msg, msg)
    meta = cs.get_meta("HSK word lists")
    check("file list stored", {f["name"] for f in meta["files"]} == {"hsk words.txt", "hsk_list.csv"})
    check("not indexed yet", meta["indexed"] is False)

    cs.create_collection("Dictionaries")
    cs.add_files("Dictionaries", [str(files["json"]), str(files["md"])])
    cs.create_collection("Examples")
    cs.add_files("Examples", [str(files["docx"]), str(files["png"]), str(files["pdf_scanned"])])
    check("three collections", set(cs.collection_names()) ==
          {"HSK word lists", "Dictionaries", "Examples"})

    msg = cs.remove_file("Dictionaries", "grammar.md")
    check("remove file", "grammar.md" not in
          {f["name"] for f in cs.get_meta("Dictionaries")["files"]})
    cs.add_files("Dictionaries", [str(files["md"])])  # put it back

    preview = cs.preview_file("HSK word lists", "hsk_list.csv")
    check("preview shows provenance", "rows" in preview and "喜欢" in preview)


def test_index_and_search():
    for name in ("HSK word lists", "Dictionaries", "Examples"):
        msg = cs.rebuild_index(name)
        check(f"index {name}", msg.startswith("✅"), msg)
        meta = cs.get_meta(name)
        check(f"meta updated {name}", meta["indexed"] and meta["chunk_count"] > 0)

    res = cs.search(["HSK word lists"], ["爱"], top_k=3, exact_word="爱")
    check("single-collection search", res and all(r["collection"] == "HSK word lists" for r in res))
    check("exact match first", res[0]["exact"] and "爱" in res[0]["text"])

    res = cs.search(["HSK word lists", "Dictionaries", "Examples"],
                    ["爱", "ài to love"], top_k=6, exact_word="爱")
    colls = {r["collection"] for r in res}
    check("multi-collection search", len(colls) >= 2, colls)
    check("provenance present", all("source" in r and "score" in r for r in res))
    line = cs.format_source_line(res[0])
    check("source line readable", res[0]["source"] in line and res[0]["collection"] in line)

    res_none = cs.search(["Dictionaries"], ["quantum physics equations"], top_k=3,
                         min_score=0.99)
    check("min_score filters", res_none == [] or all(r["exact"] for r in res_none))


if __name__ == "__main__":
    fixtures_dir = _TMP / "fixture files"
    files = make_fixtures(fixtures_dir)
    try:
        print("[extraction]")
        test_extraction(files)
        print("[collections crud]")
        test_collections_crud(files)
        print("[index + search]  (loads Qwen3-Embedding 0.6B)")
        test_index_and_search()
    finally:
        shutil.rmtree(_TMP, ignore_errors=True)
    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILURES: {FAILURES}")
        sys.exit(1)
    print("All collections tests passed.")
