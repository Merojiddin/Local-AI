"""Named document collections for the Generator tab.

Each collection is an independent folder with its own copied source files and
its own semantic index — nothing is merged into a universal index:

  data/collections/<safe_name>/
    meta.json              name, description, file list, index status
    files/                 copies of the uploaded source files
    index/embeddings.npy   one vector per chunk (mmap-read at search time)
    index/chunks.json      [{text, source, page, row, json_path, section, chunk_id}]

Extraction keeps provenance (filename, PDF page, CSV row, JSON path, MD/DOCX
section). Scanned PDF pages and images go through Apple Vision OCR. Embedding
and reranking reuse modules/documents.py (Qwen3-Embedding / Qwen3-Reranker).
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import time
from datetime import datetime
from pathlib import Path

import numpy as np

from . import documents, ocr, storage
from . import memory_manager as mm
from .safe_json import atomic_write_json, load_json, sanitize_filename

SUPPORTED_EXTS = (
    ".pdf", ".docx", ".txt", ".md", ".markdown",
    ".csv", ".json", ".png", ".jpg", ".jpeg",
)
OCR_LANGS = ["zh-Hans", "zh-Hant", "en-US", "vi"]
MIN_PDF_TEXT_CHARS = 30      # per page; below this the page is OCR'd instead
MAX_CHUNKS = 4000            # per collection
ROW_GROUP_CHARS = 500        # merge consecutive CSV rows / JSON leaves up to this
PREVIEW_CHARS = 3000


# --------------------------------------------------------------------------- #
# Collection folders
# --------------------------------------------------------------------------- #
def _coll_dir(name: str) -> Path:
    safe = sanitize_filename(name)
    return storage.collections_dir() / safe


def list_collections() -> list[dict]:
    root = storage.collections_dir()
    out = []
    if root.exists():
        for p in sorted(root.iterdir()):
            meta = load_json(p / "meta.json")
            if p.is_dir() and isinstance(meta, dict):
                out.append(meta)
    return out


def collection_names() -> list[str]:
    return [m["name"] for m in list_collections()]


def get_meta(name: str) -> dict | None:
    return load_json(_coll_dir(name) / "meta.json")


def _save_meta(name: str, meta: dict) -> None:
    meta["updated"] = datetime.now().isoformat(timespec="seconds")
    atomic_write_json(_coll_dir(name) / "meta.json", meta)


def create_collection(name: str, description: str = "") -> dict:
    name = (name or "").strip()
    if not name:
        raise ValueError("Give the collection a name.")
    target = _coll_dir(name)
    if (target / "meta.json").exists():
        raise ValueError(f"Collection '{name}' already exists.")
    (target / "files").mkdir(parents=True, exist_ok=True)
    meta = {
        "name": name,
        "description": (description or "").strip(),
        "created": datetime.now().isoformat(timespec="seconds"),
        "files": [],
        "indexed": False,
        "last_indexed": None,
        "chunk_count": 0,
        "embed_model": None,
    }
    _save_meta(name, meta)
    return meta


def delete_collection(name: str) -> str:
    target = _coll_dir(name)
    root = storage.collections_dir()
    if root not in target.parents:
        raise ValueError("Refusing to delete outside the collections folder.")
    if not (target / "meta.json").exists():
        return f"Collection '{name}' does not exist."
    shutil.rmtree(target)
    return f"🗑 Deleted collection '{name}' (models and other data untouched)."


def add_files(name: str, paths: list[str]) -> str:
    meta = get_meta(name)
    if meta is None:
        raise ValueError(f"Collection '{name}' not found — create it first.")
    files_dir = _coll_dir(name) / "files"
    files_dir.mkdir(parents=True, exist_ok=True)
    added, skipped = [], []
    for raw in paths or []:
        src = Path(raw)
        if src.suffix.lower() not in SUPPORTED_EXTS:
            skipped.append(f"{src.name} (unsupported type)")
            continue
        dest = files_dir / src.name
        shutil.copy2(src, dest)
        entry = {
            "name": src.name,
            "type": src.suffix.lower().lstrip("."),
            "size_bytes": dest.stat().st_size,
            "added": datetime.now().isoformat(timespec="seconds"),
        }
        meta["files"] = [f for f in meta["files"] if f["name"] != src.name] + [entry]
        added.append(src.name)
    if added:
        meta["indexed"] = False  # index no longer matches the file list
        meta["files"].sort(key=lambda f: f["name"])
    _save_meta(name, meta)
    msg = f"Added {len(added)} file(s) to '{name}'."
    if skipped:
        msg += " Skipped: " + ", ".join(skipped)
    if added:
        msg += " Rebuild the index to include them in retrieval."
    return msg


def remove_file(name: str, filename: str) -> str:
    meta = get_meta(name)
    if meta is None:
        raise ValueError(f"Collection '{name}' not found.")
    target = _coll_dir(name) / "files" / Path(filename).name
    if target.exists():
        target.unlink()
    before = len(meta["files"])
    meta["files"] = [f for f in meta["files"] if f["name"] != filename]
    if len(meta["files"]) != before:
        meta["indexed"] = False
    _save_meta(name, meta)
    return f"Removed '{filename}' from '{name}'. Rebuild the index to update retrieval."


def storage_usage_mb(name: str) -> float:
    return storage.dir_size_gb(_coll_dir(name)) * 1024


# --------------------------------------------------------------------------- #
# Extraction with provenance metadata
# --------------------------------------------------------------------------- #
def _record(text: str, source: str, **meta) -> dict:
    rec = {"text": text.strip(), "source": source}
    rec.update({k: v for k, v in meta.items() if v is not None})
    return rec


def _pdf_records(path: Path) -> list[dict]:
    """Per-page text via pypdf; near-empty pages are rendered and OCR'd."""
    from pypdf import PdfReader

    records = []
    reader = PdfReader(str(path))
    ocr_pages = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) >= MIN_PDF_TEXT_CHARS:
            records.append(_record(text, path.name, page=i + 1))
        else:
            ocr_pages.append(i)
    if ocr_pages:
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(str(path))
        try:
            for i in ocr_pages[: ocr.MAX_PDF_PAGES]:
                bitmap = doc[i].render(scale=2.0)
                tmp = storage.temp_dir() / f"coll_ocr_{path.stem}_{i + 1}_{int(time.time())}.png"
                bitmap.to_pil().save(tmp)
                try:
                    text = ocr.ocr_image(str(tmp), OCR_LANGS).strip()
                finally:
                    tmp.unlink(missing_ok=True)
                if text:
                    records.append(_record(text, path.name, page=i + 1, section="OCR"))
        finally:
            doc.close()
    records.sort(key=lambda r: r.get("page", 0))
    return records


def _image_records(path: Path) -> list[dict]:
    text = ocr.ocr_image(str(path), OCR_LANGS).strip()
    return [_record(text, path.name, section="OCR")] if text else []


def _docx_records(path: Path) -> list[dict]:
    import docx

    document = docx.Document(str(path))
    records, buf, section = [], [], None

    def flush():
        if buf:
            records.append(_record("\n".join(buf), path.name, section=section))
            buf.clear()

    for p in document.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading") and p.text.strip():
            flush()
            section = p.text.strip()
            buf.append(p.text.strip())
        elif p.text.strip():
            buf.append(p.text.strip())
    flush()
    return records


def _text_records(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() in (".md", ".markdown"):
        records, buf, section = [], [], None

        def flush():
            if buf and "".join(buf).strip():
                records.append(_record("\n".join(buf), path.name, section=section))
            buf.clear()

        for line in text.splitlines():
            if line.startswith("#"):
                flush()
                section = line.lstrip("# ").strip() or None
            buf.append(line)
        flush()
        return records
    return [_record(text, path.name)] if text.strip() else []


def _csv_records(path: Path) -> list[dict]:
    """Rows become 'header: value | …' lines, grouped, with row-number ranges."""
    records = []
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        rows = [r for r in reader if any((c or "").strip() for c in r)]
    if not rows:
        return []
    header = [h.strip() for h in rows[0]]
    buf, buf_start = [], None
    for n, row in enumerate(rows[1:], start=2):  # file line numbers, 1 = header
        cells = [
            f"{header[i] if i < len(header) and header[i] else f'col{i + 1}'}: {c.strip()}"
            for i, c in enumerate(row) if (c or "").strip()
        ]
        line = f"row {n}: " + " | ".join(cells)
        if buf and sum(len(b) for b in buf) + len(line) > ROW_GROUP_CHARS:
            records.append(_record("\n".join(buf), path.name,
                                   row=f"{buf_start}-{n - 1}" if n - 1 > buf_start else str(buf_start)))
            buf, buf_start = [], None
        if buf_start is None:
            buf_start = n
        buf.append(line)
    if buf:
        last = buf_start + len(buf) - 1
        records.append(_record("\n".join(buf), path.name,
                               row=f"{buf_start}-{last}" if last > buf_start else str(buf_start)))
    return records


def _json_records(path: Path) -> list[dict]:
    """Flatten to '$.path = value' lines, grouped, remembering the first path."""
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    lines: list[tuple[str, str]] = []

    def walk(node, jpath):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, f"{jpath}.{k}")
        elif isinstance(node, list):
            for i, v in enumerate(node):
                walk(v, f"{jpath}[{i}]")
        else:
            lines.append((jpath, f"{jpath} = {node}"))

    walk(data, "$")
    records, buf, buf_path = [], [], None
    for jpath, line in lines:
        if buf and sum(len(b) for b in buf) + len(line) > ROW_GROUP_CHARS:
            records.append(_record("\n".join(buf), path.name, json_path=buf_path))
            buf, buf_path = [], None
        if buf_path is None:
            buf_path = jpath
        buf.append(line)
    if buf:
        records.append(_record("\n".join(buf), path.name, json_path=buf_path))
    return records


def extract_records(path: Path) -> list[dict]:
    """Extract one file into text records with provenance metadata."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_records(path)
    if suffix in (".png", ".jpg", ".jpeg"):
        return _image_records(path)
    if suffix == ".docx":
        return _docx_records(path)
    if suffix in (".txt", ".md", ".markdown"):
        return _text_records(path)
    if suffix == ".csv":
        return _csv_records(path)
    if suffix == ".json":
        return _json_records(path)
    raise ValueError(f"Unsupported file type: {path.name}")


def _chunk_records(records: list[dict]) -> list[dict]:
    """Split long records with documents.chunk_text; metadata is inherited."""
    chunks = []
    for rec in records:
        pieces = documents.chunk_text(rec["text"]) or []
        for piece in pieces:
            c = dict(rec)
            c["text"] = piece
            chunks.append(c)
    return chunks


def preview_file(name: str, filename: str) -> str:
    path = _coll_dir(name) / "files" / Path(filename).name
    if not path.exists():
        raise ValueError(f"'{filename}' is not in collection '{name}'.")
    records = extract_records(path)
    text = "\n\n".join(
        " · ".join(
            str(v) for v in [
                r.get("source"),
                f"page {r['page']}" if r.get("page") else None,
                f"rows {r['row']}" if r.get("row") else None,
                r.get("json_path"), r.get("section"),
            ] if v
        ) + "\n" + r["text"]
        for r in records[:20]
    )
    if len(text) > PREVIEW_CHARS:
        text = text[:PREVIEW_CHARS] + "\n… (preview clipped)"
    return text or "(no text extracted)"


# --------------------------------------------------------------------------- #
# Index build / search
# --------------------------------------------------------------------------- #
def rebuild_index(name: str, progress_cb=None) -> str:
    meta = get_meta(name)
    if meta is None:
        raise ValueError(f"Collection '{name}' not found.")
    files_dir = _coll_dir(name) / "files"
    paths = [files_dir / f["name"] for f in meta["files"] if (files_dir / f["name"]).exists()]
    if not paths:
        raise ValueError(f"Collection '{name}' has no files — add some first.")

    chunks = []
    for i, p in enumerate(paths):
        if progress_cb:
            progress_cb(i / len(paths) * 0.3, f"Extracting {p.name}…")
        try:
            chunks.extend(_chunk_records(extract_records(p)))
        except Exception as exc:  # noqa: BLE001 - one bad file should not kill the build
            chunks.append(_record(f"[extraction failed: {exc}]", p.name))
        if len(chunks) > MAX_CHUNKS:
            raise ValueError(
                f"Collection '{name}' has too much text (>{MAX_CHUNKS} chunks). "
                "Split it into smaller collections."
            )
    chunks = [c for c in chunks if c["text"]]
    if not chunks:
        raise ValueError(f"No text could be extracted from the files in '{name}'.")
    for i, c in enumerate(chunks):
        c["chunk_id"] = f"{sanitize_filename(name)}#{i}"

    mm.set_task(f"Indexing '{name}': {len(chunks)} passages…")
    try:
        vectors = documents.embed_texts(
            [c["text"] for c in chunks], is_query=False,
            progress_cb=(lambda f: progress_cb(0.3 + f * 0.7, "Embedding…")) if progress_cb else None,
        )
    finally:
        mm.set_task("idle")

    index_dir = _coll_dir(name) / "index"
    index_dir.mkdir(parents=True, exist_ok=True)
    np.save(index_dir / "embeddings.npy", vectors.astype(np.float32))
    atomic_write_json(index_dir / "chunks.json", chunks, indent=None)

    from . import model_manager as mgr
    meta.update(
        indexed=True,
        last_indexed=datetime.now().isoformat(timespec="seconds"),
        chunk_count=len(chunks),
        embed_model=mgr.MODELS[documents.embed_key()]["repo"],
    )
    _save_meta(name, meta)
    return f"✅ Indexed '{name}': {len(chunks)} passages from {len(paths)} file(s)."


def _load_index(name: str):
    index_dir = _coll_dir(name) / "index"
    emb_file = index_dir / "embeddings.npy"
    if not emb_file.exists():
        return None, None
    vectors = np.load(emb_file, mmap_mode="r")  # not held fully in RAM
    chunks = load_json(index_dir / "chunks.json", default=[])
    return vectors, chunks


def search(
    names: list[str],
    query_texts: list[str],
    top_k: int = 10,
    use_rerank: bool = False,
    min_score: float = 0.0,
    exact_word: str | None = None,
    exact_first: bool = True,
    max_chars_per_chunk: int = 1200,
) -> list[dict]:
    """Search one or several collections; returns scored chunks with provenance.

    `query_texts` may hold several expanded queries — each chunk's semantic
    score is the max over all of them. When `exact_word` is set and
    `exact_first` is on, chunks containing the word verbatim are ranked first.
    """
    query_texts = [q for q in (query_texts or []) if (q or "").strip()]
    if not names or not query_texts:
        return []
    q_vecs = documents.embed_texts(query_texts, is_query=True)  # [Q, D]

    candidates = []
    for name in names:
        vectors, chunks = _load_index(name)
        if vectors is None or not chunks:
            continue
        if vectors.shape[1] != q_vecs.shape[1]:
            raise ValueError(
                f"Collection '{name}' was indexed with a different embedding "
                "model. Re-index it, or switch the embedding model back."
            )
        sims = np.asarray(vectors) @ q_vecs.T          # [N, Q]
        best = sims.max(axis=1)                        # max over expanded queries
        n_cand = min(len(chunks), max(top_k * 4, top_k) if use_rerank else top_k * 2)
        order = np.argsort(-best)[:n_cand]
        wanted = set(order.tolist())
        if exact_word:
            for i, c in enumerate(chunks):
                if exact_word in c["text"]:
                    wanted.add(i)
        for i in wanted:
            c = dict(chunks[i])
            c["collection"] = name
            c["score"] = float(best[i])
            c["exact"] = bool(exact_word) and exact_word in c["text"]
            if len(c["text"]) > max_chars_per_chunk:
                c["text"] = c["text"][:max_chars_per_chunk] + " …"
            candidates.append(c)

    if not candidates:
        return []
    candidates.sort(key=lambda c: (-c["exact"], -c["score"]) if exact_first
                    else (-c["score"], -c["exact"]))
    pool = candidates[: max(top_k * 4, top_k)]

    if use_rerank and len(pool) > 1:
        mm.set_task("Re-ranking evidence…")
        try:
            rr = documents.rerank_scores(query_texts[0], [c["text"] for c in pool])
        finally:
            mm.set_task("idle")
        for c, s in zip(pool, rr):
            c["rerank_score"] = float(s)
        pool.sort(key=lambda c: (-c["exact"], -c["rerank_score"]) if exact_first
                  else (-c["rerank_score"], -c["exact"]))
        pool = [c for c in pool if c["exact"] or c["rerank_score"] >= min_score]
    else:
        pool = [c for c in pool if c["exact"] or c["score"] >= min_score]

    return pool[: int(top_k)]


def format_source_line(c: dict) -> str:
    """'filename · collection · page 3' style provenance line for one chunk."""
    bits = [c.get("source", "?"), c.get("collection", "")]
    if c.get("page"):
        bits.append(f"page {c['page']}")
    if c.get("row"):
        bits.append(f"rows {c['row']}")
    if c.get("json_path"):
        bits.append(str(c["json_path"]))
    if c.get("section"):
        bits.append(str(c["section"]))
    return " · ".join(str(b) for b in bits if b)
